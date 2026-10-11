from lxml import etree
from PIL import Image
from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN
from pptx.oxml.ns import qn
from pptx.util import Inches, Pt

P_NS = "http://schemas.openxmlformats.org/presentationml/2006/main"
TRANS = {"fade": "<p:fade/>", "push": '<p:push dir="l"/>', "wipe": '<p:wipe dir="r"/>',
         "cover": '<p:cover dir="l"/>', "zoom": "<p:zoom/>"}


def set_transition(slide, kind="fade", speed="med"):
    el = etree.fromstring(f'<p:transition xmlns:p="{P_NS}" spd="{speed}">{TRANS.get(kind, TRANS["fade"])}</p:transition>')
    sld = slide._element
    clr = sld.find(qn("p:clrMapOvr"))
    clr.addnext(el) if clr is not None else sld.append(el)


def add_fade_entrances(slide, shape_ids):
    """Fade-in لكل عنصر بالترتيب (كل click يظهّر عنصر). اتكتبت حسب الـ schema - جرّبها في PowerPoint."""
    n = [2]
    def nid():
        n[0] += 1; return n[0]
    clicks = ""
    for sid in shape_ids:
        a, b, c, d, e = nid(), nid(), nid(), nid(), nid()
        clicks += f'''<p:par><p:cTn id="{a}" fill="hold"><p:stCondLst><p:cond delay="indefinite"/></p:stCondLst><p:childTnLst>
<p:par><p:cTn id="{b}" fill="hold"><p:stCondLst><p:cond delay="0"/></p:stCondLst><p:childTnLst>
<p:par><p:cTn id="{c}" presetID="10" presetClass="entr" presetSubtype="0" fill="hold" nodeType="clickEffect"><p:stCondLst><p:cond delay="0"/></p:stCondLst><p:childTnLst>
<p:set><p:cBhvr><p:cTn id="{d}" dur="1" fill="hold"><p:stCondLst><p:cond delay="0"/></p:stCondLst></p:cTn><p:tgtEl><p:spTgt spid="{sid}"/></p:tgtEl><p:attrNameLst><p:attrName>style.visibility</p:attrName></p:attrNameLst></p:cBhvr><p:to><p:strVal val="visible"/></p:to></p:set>
<p:animEffect transition="in" filter="fade"><p:cBhvr><p:cTn id="{e}" dur="500"/><p:tgtEl><p:spTgt spid="{sid}"/></p:tgtEl></p:cBhvr></p:animEffect>
</p:childTnLst></p:cTn></p:par></p:childTnLst></p:cTn></p:par></p:childTnLst></p:cTn></p:par>'''
    xml = f'''<p:timing xmlns:p="{P_NS}"><p:tnLst><p:par><p:cTn id="1" dur="indefinite" restart="never" nodeType="tmRoot"><p:childTnLst>
<p:seq concurrent="1" nextAc="seek"><p:cTn id="2" dur="indefinite" nodeType="mainSeq"><p:childTnLst>{clicks}</p:childTnLst></p:cTn>
<p:prevCondLst><p:cond evt="onPrev" delay="0"><p:tgtEl><p:sldTgt/></p:tgtEl></p:cond></p:prevCondLst>
<p:nextCondLst><p:cond evt="onNext" delay="0"><p:tgtEl><p:sldTgt/></p:tgtEl></p:cond></p:nextCondLst></p:seq>
</p:childTnLst></p:cTn></p:par></p:tnLst></p:timing>'''
    slide._element.append(etree.fromstring(xml))  # timing بييجي بعد transition


def _para_style(p, rtl):
    if rtl:
        p._p.get_or_add_pPr().set("rtl", "1")
        p.alignment = PP_ALIGN.RIGHT


def build(deck, path, lang="ar", animate=True):
    rtl = lang == "ar"
    prs = Presentation(); prs.slide_width, prs.slide_height = Inches(13.333), Inches(7.5)
    for s in deck["slides"]:
        sl = prs.slides.add_slide(prs.slide_layouts[6])
        sl.background.fill.solid(); sl.background.fill.fore_color.rgb = RGBColor(0x0F, 0x17, 0x2A)
        img = s.get("image")
        text_w = 6.4 if img else 12.1
        text_x = (6.3 if rtl else 0.6) if img else 0.6
        img_x = 0.6 if rtl else 6.9
        ids = []

        tb = sl.shapes.add_textbox(Inches(text_x), Inches(0.5), Inches(text_w), Inches(1.3))
        tb.text_frame.word_wrap = True
        p = tb.text_frame.paragraphs[0]; r = p.add_run(); r.text = s["title"]
        r.font.size, r.font.bold, r.font.name = Pt(36), True, "Arial"; r.font.color.rgb = RGBColor(255, 255, 255)
        _para_style(p, rtl); ids.append(tb.shape_id)

        bx = sl.shapes.add_textbox(Inches(text_x), Inches(2.0), Inches(text_w), Inches(4.8))
        tf = bx.text_frame; tf.word_wrap = True
        for i, b in enumerate(s["bullets"]):
            p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
            r = p.add_run(); r.text = "• " + b
            r.font.size, r.font.name = Pt(22), "Arial"; r.font.color.rgb = RGBColor(0xCB, 0xD5, 0xE1)
            p.space_after = Pt(14); _para_style(p, rtl)
        ids.append(bx.shape_id)

        if img:
            w, h = Image.open(img).size
            bw, bh = 5.8, 5.4
            sc = min(bw / w, bh / h)
            pw, ph = w * sc, h * sc
            pic = sl.shapes.add_picture(img, Inches(img_x + (bw - pw) / 2), Inches(1.0 + (bh - ph) / 2),
                                        Inches(pw), Inches(ph))
            ids.append(pic.shape_id)

        sl.notes_slide.notes_text_frame.text = s.get("narration", "")
        set_transition(sl, s.get("transition", "fade"))
        if animate:
            add_fade_entrances(sl, ids)
    prs.save(path)
    return path
